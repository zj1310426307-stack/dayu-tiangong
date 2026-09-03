"""HYDRO-DATA-01 parser, round-trip, validation, and OpenAPI contracts."""

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError
import pytest

from app.hydraulic.exporters import export_nwk11_subset, export_xns11_subset
from app.hydraulic.importers import parse_hydraulic_file
from app.hydraulic.processing import _submerged_interval_metrics
from app.hydraulic.schemas import (
    CoordinateReferenceSpec,
    HydraulicCrossSectionInput,
    HydraulicExchangePayload,
    HydraulicSectionPointInput,
)
from app.hydraulic.service import capabilities
from app.hydraulic.validators import validate_exchange
from app.main import app


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]


def _network_text() -> bytes:
    """Return the declared deterministic NWK11 exchange profile."""

    return b"""// HYDRO-DATA-01-NWK11
[HYDRO_NETWORK]
  NetworkCode = 'DEMO-NET'
  NetworkName = 'Demo network'
  [BRANCH]
    Code = 'R-001'
    RiverName = 'River One'
    BranchName = 'Main branch'
    FlowDirection = 'forward'
    Point = 0, 113.1000, 23.1000
    Point = 1000, 113.1100, 23.1100
  EndSect  // BRANCH
EndSect  // HYDRO_NETWORK
"""


def _section_text() -> bytes:
    """Return the declared deterministic XNS11 exchange profile."""

    return b"""// HYDRO-DATA-01-XNS11
NetworkCode = 'DEMO-NET'
NetworkName = 'Demo network'
[CROSS_SECTION]
  SectionCode = 'XS-001'
  BranchCode = 'R-001'
  Chainage = 500
  TopographyID = 'SURVEY-2026'
  Location = 113.105, 23.105
  Point = 0, 0, 12
  Point = 1, 5, 9
  Point = 2, 10, 12
EndSect  // CROSS_SECTION
"""


def test_nwk11_subset_round_trip_preserves_branch_identity_and_points() -> None:
    """A built-in NWK11 export must be readable by the matching strict parser."""

    payload, profile, status = parse_hydraulic_file("demo.nwk11", _network_text(), 4490)
    reparsed, _, _ = parse_hydraulic_file(
        "roundtrip.nwk11", export_nwk11_subset(payload), 4490
    )
    assert profile == "hydro-data-01-pfs-subset-v1"
    assert status == "ROUNDTRIP_VALIDATED_ONLY"
    assert reparsed.branches == payload.branches


def test_xns11_subset_round_trip_preserves_raw_profile() -> None:
    """A built-in XNS11 export must preserve reach, chainage, topo ID, and X/Z values."""

    payload, profile, status = parse_hydraulic_file("demo.xns11", _section_text(), 4490)
    reparsed, _, _ = parse_hydraulic_file(
        "roundtrip.xns11", export_xns11_subset(payload), 4490
    )
    assert profile == "hydro-data-01-xns11-subset-v1"
    assert status == "ROUNDTRIP_VALIDATED_ONLY"
    assert reparsed.sections == payload.sections


def test_excel_parser_accepts_bilingual_network_and_section_sheets() -> None:
    """The reviewed bilingual workbook columns normalize into one exchange DTO."""

    workbook = Workbook()
    network = workbook.active
    network.title = "河网 Network"
    network.append(
        ["网络编码", "网络名称", "河段编码", "河流名称", "河段名称", "流向", "桩号", "X坐标", "Y坐标"]
    )
    network.append(["DEMO-NET", "示例河网", "R-001", "河流一", "主河段", "forward", 0, 113.1, 23.1])
    network.append(["DEMO-NET", "示例河网", "R-001", "河流一", "主河段", "forward", 1000, 113.11, 23.11])
    sections = workbook.create_sheet("断面 Cross Section")
    sections.append(
        ["断面编号", "河段编码", "桩号", "地形编号", "点序", "距离", "高程", "断面位置X", "断面位置Y"]
    )
    for sequence, distance, elevation in ((0, 0, 12), (1, 5, 9), (2, 10, 12)):
        sections.append(
            ["XS-001", "R-001", 500, "SURVEY-2026", sequence, distance, elevation, 113.105, 23.105]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    payload, profile, _ = parse_hydraulic_file("template.xlsx", buffer.getvalue(), 4490)
    assert profile == "hydraulic-xlsx-v2"
    assert payload.network_code == "DEMO-NET"
    assert [branch.code for branch in payload.branches] == ["R-001"]
    assert [section.section_code for section in payload.sections] == ["XS-001"]


def test_excel_parser_accepts_mike11_six_column_grouped_profile_layout() -> None:
    """The user-facing MIKE11 layout may leave identity cells blank on point rows."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["ID", "TOPOID", "里程", "偏移", "高程", "river_name"])
    sheet.append([1, 11, 500, 0, 12.0, "gaominghe"])
    sheet.append([None, None, None, 5, 9.0, None])
    sheet.append([None, None, None, 10, 12.5, None])
    sheet.append([1, 12, 500, 0, 12.2, "gaominghe"])
    sheet.append([None, None, None, 5, 9.2, None])
    sheet.append([None, None, None, 10, 12.7, None])
    buffer = BytesIO()
    workbook.save(buffer)

    payload, profile, _ = parse_hydraulic_file(
        "mike11-cross-sections.xlsx", buffer.getvalue(), 4547
    )

    assert profile == "hydraulic-xlsx-v2"
    assert payload.branches == []
    assert [(section.section_code, section.topography_id) for section in payload.sections] == [
        ("1", "11"),
        ("1", "12"),
    ]
    assert all(section.branch_code == "gaominghe" for section in payload.sections)
    assert [point.distance for point in payload.sections[0].points] == [0, 5, 10]
    assert [point.sequence for point in payload.sections[0].points] == [0, 1, 2]


def test_reviewed_templates_parse_network_and_mike11_grouped_section_contracts() -> None:
    """Both reviewed workbooks must remain parser-valid after the template change."""

    root = REPOSITORY_ROOT / "outputs" / "HYDRO-DATA-01-20260818"
    network, _, _ = parse_hydraulic_file(
        "river_network.xlsx", (root / "river_network.xlsx").read_bytes(), 4547
    )
    sections, _, _ = parse_hydraulic_file(
        "cross_section.xlsx", (root / "cross_section.xlsx").read_bytes(), 4547
    )
    assert network.branches[0].source_revision == "SURVEY-2026"
    assert network.branches[0].points[0].point_code == "BP-001"
    assert network.branches[0].points[0].z == 12.3
    profile = sections.sections[0]
    assert profile.section_code == "XS-DEMO-001"
    assert profile.topography_id == "TOPO-2026"
    assert profile.branch_code == "RIVER-DEMO-001"
    assert [point.distance for point in profile.points] == [0, 5, 10]
    assert [point.marker_type for point in profile.points] == ["none", "none", "none"]
    assert profile.roughness_zones == []
    reparsed, _, _ = parse_hydraulic_file(
        "roundtrip.xns11", export_xns11_subset(sections), 4547
    )
    assert reparsed.sections == sections.sections


def test_cross_section_only_validation_uses_persisted_branch_chainage_range() -> None:
    """A later profile import must not clamp an out-of-range section to an endpoint."""

    payload = HydraulicExchangePayload(
        network_code="EXISTING-NET",
        network_name="Existing network",
        source_srid=4547,
        source_kind="excel",
        branches=[],
        sections=[HydraulicCrossSectionInput(
            section_code="DM22",
            branch_code="gaominghe",
            chainage=5725.8932,
            topography_id="11",
            points=[
                HydraulicSectionPointInput(sequence=0, distance=0, elevation=10),
                HydraulicSectionPointInput(sequence=1, distance=5, elevation=8),
                HydraulicSectionPointInput(sequence=2, distance=10, elevation=10),
            ],
        )],
    )

    issues = validate_exchange(
        payload,
        known_branch_codes={"gaominghe"},
        known_branch_ranges={"gaominghe": (0.0, 5432.1266)},
    )

    outside = [issue for issue in issues if issue.code == "SECTION_CHAINAGE_OUTSIDE_BRANCH"]
    assert len(outside) == 1
    assert outside[0].entity_ref == "DM22"
    assert outside[0].context == {
        "chainage": 5725.8932,
        "start": 0.0,
        "end": 5432.1266,
    }


def test_cross_section_groups_must_follow_upstream_to_downstream_order() -> None:
    """Section groups on one branch must be entered in nondecreasing chainage order."""

    points = [
        HydraulicSectionPointInput(sequence=0, distance=0, elevation=10),
        HydraulicSectionPointInput(sequence=1, distance=5, elevation=8),
        HydraulicSectionPointInput(sequence=2, distance=10, elevation=10),
    ]
    payload = HydraulicExchangePayload(
        network_code="ORDER-NET",
        network_name="Section order",
        source_srid=4547,
        source_kind="excel",
        branches=[],
        sections=[
            HydraulicCrossSectionInput(
                section_code="DM02",
                branch_code="gaominghe",
                chainage=500,
                topography_id="11",
                points=points,
            ),
            HydraulicCrossSectionInput(
                section_code="DM01",
                branch_code="gaominghe",
                chainage=100,
                topography_id="11",
                points=points,
            ),
        ],
    )

    issues = validate_exchange(
        payload,
        known_branch_codes={"gaominghe"},
        known_branch_ranges={"gaominghe": (0.0, 1000.0)},
    )

    order = [issue for issue in issues if issue.code == "SECTION_CHAINAGE_ORDER_INVALID"]
    assert len(order) == 1
    assert order[0].entity_ref == "DM01"
    assert order[0].context == {
        "branch_code": "gaominghe",
        "previous_section_code": "DM02",
        "previous_chainage": 500.0,
        "chainage": 100.0,
    }


def test_cross_section_profiles_may_share_one_chainage() -> None:
    """Multiple TOPOID profiles at one section location keep a valid table order."""

    points = [
        HydraulicSectionPointInput(sequence=0, distance=0, elevation=10),
        HydraulicSectionPointInput(sequence=1, distance=5, elevation=8),
        HydraulicSectionPointInput(sequence=2, distance=10, elevation=10),
    ]
    payload = HydraulicExchangePayload(
        network_code="ORDER-NET",
        network_name="Section order",
        source_srid=4547,
        source_kind="excel",
        branches=[],
        sections=[
            HydraulicCrossSectionInput(
                section_code="DM01",
                branch_code="gaominghe",
                chainage=100,
                topography_id="11",
                points=points,
            ),
            HydraulicCrossSectionInput(
                section_code="DM01",
                branch_code="gaominghe",
                chainage=100,
                topography_id="12",
                points=points,
            ),
        ],
    )

    issues = validate_exchange(
        payload,
        known_branch_codes={"gaominghe"},
        known_branch_ranges={"gaominghe": (0.0, 1000.0)},
    )

    assert not any(issue.code == "SECTION_CHAINAGE_ORDER_INVALID" for issue in issues)


def test_segmented_roughness_geometry_integrates_each_wetted_interval() -> None:
    """Compound-section conveyance inputs must be integrated by roughness interval."""

    points = [(0.0, 2.0), (5.0, 0.0), (10.0, 2.0)]
    left = _submerged_interval_metrics(points, 1.0, 0.0, 5.0)
    right = _submerged_interval_metrics(points, 1.0, 5.0, 10.0)
    assert left == pytest.approx((1.25, 2.5, (2.5 ** 2 + 1.0) ** 0.5))
    assert right == pytest.approx(left)


def test_csv_parser_keeps_topography_versions_and_roughness_zones_separate() -> None:
    """CSV exchange rows may carry multiple survey profiles for one section location."""

    header = (
        "record_type,network_code,network_name,section_code,section_name,branch_code,"
        "chainage,topography_id,survey_date,survey_method,default_manning_n,"
        "location_x,location_y,axis_x,axis_y,sequence,distance,elevation,marker_type,"
        "roughness_zone_order,roughness_start,roughness_end,roughness_n,roughness_type\n"
    )
    rows: list[str] = []
    for topo, base in (("SURVEY-A", 10), ("SURVEY-B", 11)):
        values = (
            (0, 0, base + 2, 0, 4),
            (1, 5, base, 4, 7),
            (2, 10, base + 2, 7, 10),
        )
        for sequence, distance, elevation, zone_start, zone_end in values:
            rows.append(
                f"section_point,CSV-NET,CSV Network,XS-001,Section 1,R-001,500,{topo},"
                f"2026-08-18,RTK,0.03,500000,2550000,{500000 + sequence},2550000,"
                f"{sequence},{distance},{elevation},none,{sequence},{zone_start},"
                f"{zone_end},0.03,custom"
            )
    payload, profile, _ = parse_hydraulic_file(
        "profiles.csv", (header + "\n".join(rows)).encode(), 4547
    )
    assert profile == "hydraulic-csv-v1"
    assert payload.network_code == "CSV-NET"
    assert [(item.section_code, item.topography_id) for item in payload.sections] == [
        ("XS-001", "SURVEY-A"), ("XS-001", "SURVEY-B")
    ]
    assert all(len(item.roughness_zones) == 3 for item in payload.sections)


def test_coordinate_contract_rejects_geographic_engineering_or_implicit_units() -> None:
    """No topology workflow may silently treat angular coordinates as metres."""

    with pytest.raises(ValidationError):
        CoordinateReferenceSpec(
            source_crs="EPSG:4490", engineering_crs="EPSG:4490",
            coordinate_mode="geographic", axis_mapping="x_easting_y_northing",
            horizontal_unit="degree", vertical_datum="1985 National Height Datum",
        )
    with pytest.raises(ValidationError):
        CoordinateReferenceSpec(
            source_crs="EPSG:4547", engineering_crs="EPSG:4547",
            coordinate_mode="projected", axis_mapping="x_easting_y_northing",
            horizontal_unit="degree", vertical_datum="1985 National Height Datum",
        )


def test_validation_fails_closed_for_unknown_section_branch() -> None:
    """A section cannot be committed unless its referenced branch is resolvable."""

    payload = HydraulicExchangePayload(
        network_code="SECTIONS",
        network_name="Sections only",
        source_srid=4490,
        source_kind="api",
        sections=[
            HydraulicCrossSectionInput(
                section_code="XS-MISSING",
                branch_code="UNKNOWN",
                chainage=10,
                points=[
                    HydraulicSectionPointInput(sequence=0, distance=0, elevation=2),
                    HydraulicSectionPointInput(sequence=1, distance=5, elevation=1),
                    HydraulicSectionPointInput(sequence=2, distance=10, elevation=2),
                ],
            )
        ],
    )
    issues = validate_exchange(payload)
    assert any(issue.code == "SECTION_BRANCH_MISSING" and issue.severity == "error" for issue in issues)


def test_openapi_exposes_complete_hydraulic_management_surface() -> None:
    """The generated client source must be able to discover every management operation."""

    paths = TestClient(app).get("/openapi.json").json()["paths"]
    required = {
        "/api/v1/hydraulic/capabilities",
        "/api/v1/hydraulic/engine-capabilities",
        "/api/v1/hydraulic/networks",
        "/api/v1/hydraulic/networks/{network_id}/graph",
        "/api/v1/hydraulic/structures",
        "/api/v1/hydraulic/structures/{structure_id}",
        "/api/v1/hydraulic/structures/{structure_id}/scenarios/{case_id}",
        "/api/v1/hydraulic/cross-sections/{section_id}",
        "/api/v1/hydraulic/imports",
        "/api/v1/hydraulic/imports/preview",
        "/api/v1/hydraulic/imports/commit",
        "/api/v1/hydraulic/networks/{network_id}/topology",
        "/api/v1/hydraulic/branches/{branch_id}/reverse",
        "/api/v1/hydraulic/branches/{branch_id}/recalculate-chainage",
        "/api/v1/hydraulic/cross-sections/{section_id}/locate",
        "/api/v1/hydraulic/profiles/{profile_id}/process",
        "/api/v1/hydraulic/profiles/process-batch",
        "/api/v1/hydraulic/validation/run",
        "/api/v1/hydraulic/validation/{run_code}",
        "/api/v1/hydraulic/exports/network.nwk11",
        "/api/v1/hydraulic/exports/cross-sections.xns11",
        "/api/v1/hydraulic/templates/{template_name}",
        "/api/v1/model/readiness",
        "/api/v1/model/preview",
        "/api/v1/model/tasks",
        "/api/v1/model/results/{task_id}",
    }
    assert required <= paths.keys()


def test_capability_contract_never_claims_native_nwk11_validation() -> None:
    """The API must distinguish round-trip subset support from licensed native validation."""

    result = capabilities()
    assert result.native_nwk11_available is False
    assert "external acceptance step" in result.limitation


def test_reviewed_excel_templates_are_downloadable_workbooks() -> None:
    """Both public template names must resolve to actual XLSX packages."""

    client = TestClient(app)
    for template_name in ("river-network", "cross-section"):
        response = client.get(f"/api/v1/hydraulic/templates/{template_name}")
        assert response.status_code == 200
        assert response.content.startswith(b"PK")
        assert "spreadsheetml.sheet" in response.headers["content-type"]
